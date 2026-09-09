from docx import Document
from app.utils.exceptions import ResumeParserError

def extract_docx_text(path):
    try:
        doc = Document(path)
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(chunks).strip()
        if not text:
            raise ResumeParserError("NO_TEXT_FOUND", "No readable text was found in the DOCX.")
        return text
    except ResumeParserError:
        raise
    except Exception as exc:
        raise ResumeParserError("CORRUPTED_DOCX", "Unable to read the DOCX file.") from exc
